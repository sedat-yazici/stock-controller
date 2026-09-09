import os
import sys
import time
import random
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter
import requests
from dotenv import load_dotenv

# Import core crawler routines
import crawler_core

# Ensure UTF-8 output on Windows terminal
if sys.stdout is not None:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Load environment variables from the current working directory
env_path = os.path.join(os.getcwd(), ".env")
if not os.path.exists(env_path):
    print("Error: .env file does not exist in the current folder.")
    sys.exit(1)

load_dotenv(dotenv_path=env_path, override=True)

LOGIN_URL = os.getenv("LOGIN_URL")
USERNAME = os.getenv("USERNAME")
PASSWORD = os.getenv("PASSWORD")

# Cache file in the current working directory (will be created automatically on save)
CACHE_FILE = os.path.join(os.getcwd(), "session_cache.json")

def load_excel_safely(file_path):
    """Loads Excel workbook safely, handling potential file lock issues."""
    try:
        return openpyxl.load_workbook(file_path)
    except PermissionError:
        print(f"Warning: {file_path} appears to be locked. Attempting to copy and read...")
        temp_path = file_path.replace(".xlsx", "_temp_read.xlsx")
        cmd = f'powershell -Command "Copy-Item \'{file_path}\' \'{temp_path}\' -Force"'
        os.system(cmd)
        try:
            wb = openpyxl.load_workbook(temp_path)
            return wb
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)
    except Exception as e:
        print(f"Error reading Excel file: {e}")
        sys.exit(1)

def save_excel_safely(wb, output_path):
    """Saves workbook safely, handling potential permission issues without altering styles."""
    try:
        wb.save(output_path)
        print(f"Successfully saved results to: {output_path}")
    except PermissionError:
        print(f"Error: Permission denied when writing to {output_path}. Is the file open in Excel?")
        alternative_path = output_path.replace(".xlsx", f"_backup_{int(time.time())}.xlsx")
        wb.save(alternative_path)
        print(f"Saved to alternative path: {alternative_path}")
    except Exception as e:
        print(f"Critical error saving Excel file: {e}")

def main():
    if not LOGIN_URL or not USERNAME or not PASSWORD:
        print("Error: Missing credentials in environment variables (.env file).")
        sys.exit(1)
        
    input_file = os.getenv("INPUT_EXCEL_PATH", r"input.xlsx")
    output_file = os.getenv("OUTPUT_EXCEL_PATH", r"results.xlsx")
    
    print(f"Loading input file: {input_file}")
    wb = load_excel_safely(input_file)
    sheet = wb.active
    
    # Ensure sheet views always display grid lines
    sheet.sheet_view.showGridLines = True
    if hasattr(sheet.views, 'sheetView') and sheet.views.sheetView:
        for sv in sheet.views.sheetView:
            sv.showGridLines = True
            
    # Read headers
    header_row = []
    for col_idx in range(1, sheet.max_column + 1):
        header_row.append(sheet.cell(row=1, column=col_idx).value)
        
    code_col_idx = None
    for idx, val in enumerate(header_row):
        if val and str(val).lower().replace(" ", "") in ["ürünkodu", "urunkodu", "code", "stokkodu"]:
            code_col_idx = idx + 1
            break
            
    if not code_col_idx:
        cols_str = ", ".join(map(str, [v for v in header_row if v is not None]))
        print(f"Hata: ürün kodu sütunu bulunamadı. Excel dosyasındaki sütunlar: {cols_str}")
        sys.exit(1)
        
    ref_header_cell = sheet.cell(row=1, column=code_col_idx)
    
    # Target columns
    price_col_idx = None
    qty_col_idx = None
    
    for idx, val in enumerate(header_row):
        if val == 'Alış Fiyatı':
            price_col_idx = idx + 1
        elif val == 'Ürün Adedi':
            qty_col_idx = idx + 1
            
    if not price_col_idx:
        price_col_idx = sheet.max_column + 1
        sheet.cell(row=1, column=price_col_idx, value="Alış Fiyatı")
        header_row.append("Alış Fiyatı")
        
    if not qty_col_idx:
        qty_col_idx = sheet.max_column + 1
        sheet.cell(row=1, column=qty_col_idx, value="Ürün Adedi")
        header_row.append("Ürün Adedi")
        
    # Apply format and column widths to result columns
    for c_idx, title in [(price_col_idx, "Alış Fiyatı"), (qty_col_idx, "Ürün Adedi")]:
        h_cell = sheet.cell(row=1, column=c_idx)
        h_cell.value = title
        if ref_header_cell.has_style:
            if ref_header_cell.font: h_cell.font = copy(ref_header_cell.font)
            if ref_header_cell.border: h_cell.border = copy(ref_header_cell.border)
            if ref_header_cell.fill: h_cell.fill = copy(ref_header_cell.fill)
            if ref_header_cell.alignment: h_cell.alignment = copy(ref_header_cell.alignment)
        col_letter = get_column_letter(c_idx)
        curr_width = sheet.column_dimensions[col_letter].width if col_letter in sheet.column_dimensions else None
        if not curr_width or curr_width < 14:
            sheet.column_dimensions[col_letter].width = 16.0

    session = requests.Session()
    if not crawler_core.login(session, USERNAME, PASSWORD, CACHE_FILE):
        print("Aborting: Authentication failed.")
        sys.exit(1)
        
    total_rows = sheet.max_row - 1
    if total_rows <= 0:
        print("No data rows found to process.")
        sys.exit(1)
        
    print(f"Processing {total_rows} products...")
    
    for row_num in range(2, sheet.max_row + 1):
        cell_val = sheet.cell(row=row_num, column=code_col_idx).value
        product_code = str(cell_val).strip() if cell_val is not None else ""
        
        if not product_code or product_code.lower() in ["nan", "none", ""]:
            continue
            
        print(f"[{row_num - 1}/{total_rows}] Searching for product code: {product_code}")
        
        price, stock = crawler_core.search_and_extract(session, product_code)
        
        price_cell = sheet.cell(row=row_num, column=price_col_idx, value=price)
        qty_cell = sheet.cell(row=row_num, column=qty_col_idx, value=stock)
        
        # Preserve cell layout and grid borders from reference row
        ref_data_cell = sheet.cell(row=row_num, column=code_col_idx)
        if ref_data_cell.has_style:
            for target_cell in (price_cell, qty_cell):
                if ref_data_cell.border:
                    target_cell.border = copy(ref_data_cell.border)
                if ref_data_cell.font and not target_cell.font:
                    target_cell.font = copy(ref_data_cell.font)
                if ref_data_cell.fill and (not target_cell.fill or not target_cell.fill.fill_type):
                    target_cell.fill = copy(ref_data_cell.fill)
                    
        print(f"  -> Extracted Price: {price}, Stock: {stock}")
        
        # Mimic human delay
        if row_num - 1 < total_rows:
            sleep_time = random.uniform(2.0, 5.0)
            print(f"  -> Waiting {sleep_time:.2f} seconds before next search...")
            time.sleep(sleep_time)
            
    # Ensure all data rows preserve the table grid layout and borders even if skipped
    for r_num in range(2, sheet.max_row + 1):
        ref_r = sheet.cell(row=r_num, column=code_col_idx)
        if ref_r.has_style and ref_r.border and ref_r.border.left and ref_r.border.left.style:
            for c_idx in (price_col_idx, qty_col_idx):
                target_c = sheet.cell(row=r_num, column=c_idx)
                if not target_c.border or not target_c.border.left or not target_c.border.left.style:
                    target_c.border = copy(ref_r.border)
                if ref_r.font and (not target_c.font or not target_c.font.name):
                    target_c.font = copy(ref_r.font)

    # Ensure grid lines remain enabled
    sheet.sheet_view.showGridLines = True
    if hasattr(sheet.views, 'sheetView') and sheet.views.sheetView:
        for sv in sheet.views.sheetView:
            sv.showGridLines = True

    # Save the updated copy with all layout and styles preserved
    save_excel_safely(wb, output_file)

if __name__ == "__main__":
    main()
